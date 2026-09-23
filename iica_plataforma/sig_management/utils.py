# tu_app/utils.py
import re
from docx import Document

def extraer_etiquetas_docx(archivo_docx):
    """Lee un archivo docx y extrae todas las etiquetas únicas tipo {{ etiqueta }}."""
    doc = Document(archivo_docx)
    texto_total = ""
    
    # Extraer texto de párrafos
    for p in doc.paragraphs:
        texto_total += p.text + " "
        
    # Extraer texto de tablas (por si las etiquetas están en tablas)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texto_total += cell.text + " "
                
    # Buscar patrones entre llaves dobles
    etiquetas = re.findall(r'\{\{\s*([\w_]+)\s*\}\}', texto_total)
    return list(set(etiquetas)) # Devuelve lista de etiquetas únicas sin duplicados